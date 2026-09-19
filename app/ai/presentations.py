"""Presentation generation: AI builds a JSON slide outline, then we render
it either as JSON (web slides) or a real .pptx via python-pptx."""

import json
import re
import uuid
from io import BytesIO

from app.ai import client

OUTLINE_PROMPT = (
    "You are an expert presentation designer. Create a slide deck outline "
    "for the given topic. Respond with ONLY valid JSON (no markdown fences, "
    "no commentary) matching exactly this schema:\n"
    "{{\n"
    '  "title": "deck title",\n'
    '  "subtitle": "one-line subtitle",\n'
    '  "slides": [\n'
    '    {{"title": "slide title", "bullets": ["point 1", "point 2", "..."],\n'
    '      "notes": "1-3 sentences of speaker notes"}}\n'
    "  ]\n"
    "}}\n"
    "Rules: exactly {num_slides} slides; 3-5 concise bullets per slide "
    "(each under 90 characters); include a strong opening and closing slide; "
    "audience: {audience}; tone: {tone}; language: {language}."
)


def _extract_json(text: str) -> dict:
    text = text.strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.MULTILINE)
    start, end = text.find("{"), text.rfind("}")
    if start == -1 or end == -1:
        raise ValueError("AI did not return JSON")
    data = json.loads(text[start:end + 1])
    if not isinstance(data.get("slides"), list) or not data["slides"]:
        raise ValueError("outline missing slides")
    slides = []
    for s in data["slides"][:30]:
        if not isinstance(s, dict):
            continue
        bullets = [str(b)[:160] for b in (s.get("bullets") or []) if str(b).strip()]
        slides.append({
            "title": str(s.get("title") or "")[:120],
            "bullets": bullets[:6],
            "notes": str(s.get("notes") or "")[:600],
        })
    if not slides:
        raise ValueError("outline has no usable slides")
    data["slides"] = slides
    data["title"] = str(data.get("title") or "Presentation")[:150]
    data["subtitle"] = str(data.get("subtitle") or "")[:200]
    data["id"] = str(uuid.uuid4())
    return data


async def generate_outline(
    topic: str,
    *,
    num_slides: int = 8,
    audience: str = "general audience",
    tone: str = "professional but engaging",
    language: str = "English",
) -> dict:
    prompt = OUTLINE_PROMPT.format(
        num_slides=max(3, min(num_slides, 20)),
        audience=audience[:60],
        tone=tone[:60],
        language=language[:30],
    )
    raw = await client.complete(
        [{"role": "system", "content": prompt},
         {"role": "user", "content": topic[:2000]}],
        temperature=0.6,
    )
    try:
        return _extract_json(raw)
    except (ValueError, json.JSONDecodeError):
        # One repair attempt with a stricter instruction
        raw2 = await client.complete(
            [{"role": "system",
              "content": "Output ONLY minified valid JSON. No prose."},
             {"role": "user", "content":
                 "Convert this to valid JSON per the schema. Output JSON only:\n"
                 + raw[:6000]}],
            temperature=0.0,
        )
        return _extract_json(raw2)


# ----------------------------------------------------------------- PPTX

THEMES = {
    "ocean": ("0B1F3A", "3EC1D3", "FFFFFF"),
    "ember": ("1A0F0B", "FF7A45", "FFF6F0"),
    "forest": ("0E2A1A", "6FCF97", "F2FFF7"),
    "violet": ("1D1030", "A78BFA", "F5F0FF"),
}


def outline_to_pptx(outline: dict, theme: str = "ocean") -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    bg_hex, accent_hex, text_hex = THEMES.get(theme, THEMES["ocean"])
    bg, accent, fg = (RGBColor.from_string(c) for c in (bg_hex, accent_hex, text_hex))

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def _style(slide, dark_text=False):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = bg

    # ---- Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    _style(slide)
    tx = slide.shapes.add_textbox(Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.8))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = outline.get("title") or "Presentation"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = fg
    p.alignment = PP_ALIGN.LEFT
    if outline.get("subtitle"):
        p2 = tf.add_paragraph()
        p2.text = outline["subtitle"]
        p2.font.size = Pt(20)
        p2.font.color.rgb = accent
    bar = slide.shapes.add_shape(1, Inches(0.95), Inches(2.1), Inches(2.2), Pt(4))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()

    # ---- Content slides
    for s in outline.get("slides", []):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _style(slide)
        # accent bar
        bar = slide.shapes.add_shape(1, Inches(0.95), Inches(0.75), Inches(1.4), Pt(4))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.fill.background()
        # title
        tx = slide.shapes.add_textbox(Inches(0.9), Inches(0.95), Inches(11.5), Inches(1.1))
        tf = tx.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = s.get("title") or ""
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = fg
        # bullets
        bullets = s.get("bullets") or []
        if bullets:
            bx = slide.shapes.add_textbox(Inches(1.1), Inches(2.3), Inches(11.0), Inches(4.4))
            btf = bx.text_frame
            btf.word_wrap = True
            for i, b in enumerate(bullets):
                para = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
                para.text = "•  " + b
                para.font.size = Pt(20)
                para.font.color.rgb = fg
                para.space_after = Pt(12)
        # speaker notes
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]

    # ---- Closing slide
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _style(slide)
    tx = slide.shapes.add_textbox(Inches(0.9), Inches(2.8), Inches(11.5), Inches(1.5))
    p = tx.text_frame.paragraphs[0]
    p.text = "Thank you"
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = accent

    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()
