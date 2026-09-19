"""Keyless end-to-end check of a *deployed* Aether instance.

Runs against a real URL with ZERO API keys configured, exactly like a normal
user: registers a throwaway account, streams a chat answer, builds a
presentation, exports it as .pptx and HTML, generates an image, and asks for
TTS audio. Everything is measured and printed as a pass/fail report.

    python tests/e2e_preview.py --base-url https://ai-xxxx.vercel.app

Exit code 0 = everything that is expected to work, worked.
"""

from __future__ import annotations

import argparse
import io
import json
import os
import re
import secrets
import sys
import time
import zipfile
from dataclasses import dataclass, field

import httpx

TIMEOUT = httpx.Timeout(180.0, connect=20.0)


@dataclass
class Results:
    checks: list[tuple[str, bool, str]] = field(default_factory=list)
    warnings: list[tuple[str, str]] = field(default_factory=list)

    def add(self, name: str, ok: bool, detail: str = "") -> bool:
        mark = "PASS" if ok else "FAIL"
        print(f"[{mark}] {name}{(' — ' + detail) if detail else ''}")
        self.checks.append((name, ok, detail))
        return ok

    def skip(self, name: str, detail: str = "") -> None:
        print(f"[SKIP] {name}{(' — ' + detail) if detail else ''}")

    def note(self, text: str) -> None:
        print(f"      {text}")

    def warn(self, name: str, detail: str = "") -> None:
        """Soft check: informational only, never fails the run."""
        print(f"[WARN] {name}{(' — ' + detail) if detail else ''}")
        self.warnings.append((name, detail))

    @property
    def failed(self) -> list[tuple[str, bool, str]]:
        return [c for c in self.checks if not c[1]]

    def summary_markdown(self) -> str:
        lines = ["| Check | Result | Detail |", "|---|---|---|"]
        for name, ok, detail in self.checks:
            lines.append(f"| {name} | {'✅' if ok else '❌'} | {detail[:160]} |")
        total = len(self.checks)
        good = total - len(self.failed)
        out = f"**{good}/{total} checks passed**\n\n" + "\n".join(lines)
        if self.warnings:
            out += "\n\n**Soft checks (informational — never fail the run)**\n\n"
            out += "| Check | Detail |\n|---|---|\n"
            out += "\n".join(f"| {n} | {d[:160]} |" for n, d in self.warnings)
        return out


R = Results()


BYPASS = os.getenv("AETHER_BYPASS", "").strip()


def client() -> httpx.Client:
    headers = {"User-Agent": "aether-e2e"}
    if BYPASS:
        # Vercel deployment protection (preview URLs): "Protection Bypass for
        # Automation" secret lets CI reach the preview without SSO.
        headers["x-vercel-protection-bypass"] = BYPASS
        headers["x-vercel-set-bypass-cookie"] = "true"
    return httpx.Client(timeout=TIMEOUT, follow_redirects=True, headers=headers)


def read_sse(resp: httpx.Response, on_event, deadline: float) -> str:
    """Consume an SSE response, calling on_event({...}) per event.
    Returns the raw accumulated text for debugging."""
    raw = []
    buf = ""
    for chunk in resp.iter_text():
        raw.append(chunk)
        buf += chunk
        while "\n\n" in buf:
            frame, buf = buf.split("\n\n", 1)
            payload = None
            for line in frame.splitlines():
                if line.startswith("data:"):
                    payload = line[5:].strip()
            if payload is None or payload == "[DONE]":
                continue
            try:
                on_event(json.loads(payload))
            except json.JSONDecodeError:
                continue
        if time.time() > deadline:
            break
    return "".join(raw)


# --------------------------------------------------------------------- tests

def test_static(c: httpx.Client, base: str) -> None:
    t0 = time.time()
    r = c.get(f"{base}/api/health")
    R.add("health endpoint", r.status_code == 200,
          f"{r.status_code} in {time.time() - t0:.2f}s")

    r = c.get(f"{base}/")
    html = r.text
    looks_like_app = "<title>Aether</title>" in html
    R.add("app shell loads", r.status_code == 200 and looks_like_app,
          f"{len(html)} bytes")
    if not looks_like_app:
        R.add("deployment reachable without SSO", False,
              "got an HTML login/protection page — is this a protected preview?")

    logo = c.get(f"{base}/icons/icon-512.png")
    R.add("logo asset served", logo.status_code == 200 and len(logo.content) > 2000,
          f"{len(logo.content)} bytes")

    manifest = c.get(f"{base}/manifest.json")
    R.add("manifest served", manifest.status_code == 200)

    # Every asset the shell references must exist: scrape index.html instead of
    # hard-coding filenames, so a refactor cannot leave a dead link behind.
    import re
    refs = sorted(set(re.findall(r'(?:src|href)="(/assets/[^"]+|/icons/[^"]+)"', html)))
    missing = []
    total = 0
    for ref in refs:
        asset = c.get(f"{base}{ref}")
        total += len(asset.content)
        if asset.status_code != 200 or not asset.content:
            missing.append(f"{ref} ({asset.status_code})")
    R.add(f"every shell asset served ({len(refs)} files)", not missing,
          ", ".join(missing) if missing else f"{total // 1024} KB total")
    for name in ("core.js", "chat.js", "voice.js", "slides.js", "images.js",
                 "settings.js", "tools.js", "admin.js", "app.js"):
        script = c.get(f"{base}/assets/js/{name}")
        R.add(f"module {name} served", script.status_code == 200 and len(script.content) > 800,
              f"{len(script.content)} bytes")


def test_auth(c: httpx.Client, base: str) -> str | None:
    email = f"e2e-{secrets.token_hex(6)}@example.com"
    password = "e2e-Passw0rd-" + secrets.token_hex(4)
    r = c.post(f"{base}/api/auth/register",
               json={"email": email, "password": password})
    if r.status_code not in (200, 201):
        R.add("register", False, f"{r.status_code}: {r.text[:160]}")
        return None
    token = (r.json() or {}).get("access_token")
    R.add("register + token", bool(token), email)
    if not token:
        return None
    c.headers["Authorization"] = f"Bearer {token}"
    me = c.get(f"{base}/api/auth/me")
    R.add("auth/me", me.status_code == 200 and me.json().get("email") == email)
    return token


def test_chat(c: httpx.Client, base: str, conversation: str) -> str | None:
    """Stream one short answer; returns the answer text."""
    events: list[dict] = []
    answer = ""
    stream_error = ""
    first_delta_at: float | None = None
    t0 = time.time()

    with c.stream("POST", f"{base}/api/chat/stream",
                  json={"conversation_id": conversation,
                        "message": "In one short sentence: what is 2+2?",
                        "mode": "ai"},
                  headers={"Accept": "text/event-stream"}) as resp:
        if resp.status_code != 200:
            body = resp.read().decode(errors="replace")[:200]
            R.add("chat stream HTTP", False, f"{resp.status_code}: {body}")
            return None
        R.add("chat stream HTTP", True, f"{resp.status_code}")

        def on_event(ev: dict) -> None:
            nonlocal answer, first_delta_at, stream_error
            events.append(ev)
            if ev.get("type") == "delta":
                if first_delta_at is None:
                    first_delta_at = time.time() - t0
                answer += ev.get("text", "")
            elif ev.get("type") == "final":
                answer = ev.get("text") or answer
            elif ev.get("type") == "error":
                stream_error = str(ev.get("detail"))[:200]

        read_sse(resp, on_event, deadline=time.time() + 150)

    types = [e.get("type") for e in events]
    R.add("chat phase + deltas", "delta" in types,
          f"events={sorted(set(types))} ttft={first_delta_at and round(first_delta_at, 2)}s")
    R.add("chat final answer", bool(answer.strip()),
          f"{len(answer)} chars" + (f" — stream error: {stream_error}" if stream_error else ""))
    if first_delta_at:
        R.add("chat first token under 30s", first_delta_at < 30,
              f"{first_delta_at:.1f}s")
    return answer.strip() or None


def test_chat_controls(c: httpx.Client, base: str, conversation: str) -> None:
    msgs_before = c.get(f"{base}/api/conversations/{conversation}").json().get("messages", [])
    n_before = len(msgs_before)
    # Regenerate must re-answer WITHOUT storing another user message.
    answer = ""
    with c.stream("POST", f"{base}/api/chat/stream",
                  json={"conversation_id": conversation, "message": "",
                        "regenerate": True, "mode": "ai"}) as resp:
        if resp.status_code != 200:
            R.add("regenerate accepted", False, f"{resp.status_code}: {resp.read().decode()[:160]}")
            return
        def on_event(ev):
            nonlocal answer
            if ev.get("type") in ("delta", "final"):
                answer += ev.get("text", "") if ev.get("type") == "delta" else ""
            if ev.get("type") == "final":
                answer = ev.get("text") or answer
        read_sse(resp, on_event, deadline=time.time() + 150)
    msgs_after = c.get(f"{base}/api/conversations/{conversation}").json().get("messages", [])
    user_msgs_before = sum(1 for m in msgs_before if m["role"] == "user")
    user_msgs_after = sum(1 for m in msgs_after if m["role"] == "user")
    R.add("regenerate keeps user messages unchanged",
          user_msgs_after == user_msgs_before,
          f"before={user_msgs_before} after={user_msgs_after}")
    R.add("regenerate produced an answer", bool(answer.strip()), f"{len(answer)} chars")


def test_teach_and_memory(c: httpx.Client, base: str) -> None:
    r = c.post(f"{base}/api/settings/memory",
               json={"content": "The user goes by Sam and prefers short answers."})
    if r.status_code == 404:
        R.skip("teach → memory", "endpoint not deployed yet")
        return
    R.add("teach saves memory", r.status_code in (200, 201), f"{r.status_code}")
    mem = c.get(f"{base}/api/settings/memory").json()
    contents = [m["content"] for m in mem.get("memories", [])]
    R.add("memory readable in settings", any("Sam" in x for x in contents),
          f"{len(contents)} memories")
    mid = next((m["id"] for m in mem["memories"] if "Sam" in m["content"]), None)
    if mid:
        upd = c.patch(f"{base}/api/settings/memory/{mid}",
                      json={"content": "The user goes by Sam and likes bullet points."})
        R.add("memory edit", upd.status_code == 200, f"{upd.status_code}")
        dele = c.delete(f"{base}/api/settings/memory/{mid}")
        R.add("memory delete", dele.status_code == 200, f"{dele.status_code}")


def test_usage(c: httpx.Client, base: str) -> None:
    r = c.get(f"{base}/api/usage")
    if r.status_code != 200:
        R.add("usage meter API", False, f"{r.status_code}")
        return
    data = r.json()
    ok = {"used_today", "limit"} <= set(data)
    R.add("usage meter API", ok, json.dumps(data)[:160])
    R.add("usage counter is counting", data.get("used_today", 0) > 0,
          f"used_today={data.get('used_today')} limit={data.get('limit')}")


def test_presentations(c: httpx.Client, base: str) -> None:
    topic = "How bees help farms"
    plan = c.post(f"{base}/api/presentations/plan",
                  json={"topic": topic, "num_slides": 6})
    if plan.status_code == 404:
        R.skip("incremental presentation plan", "endpoint not deployed yet")
        outline = None
    elif plan.status_code != 200:
        R.add("presentation plan", False, f"{plan.status_code}: {plan.text[:200]}")
        outline = None
    else:
        plan_data = plan.json()
        deck_id = plan_data.get("deck_id")
        titles = plan_data.get("slides") or []
        R.add("presentation plan", bool(deck_id and titles),
              f"{len(titles)} slide titles")
        slides = []
        t0 = time.time()
        for i in range(len(titles)):
            sr = c.post(f"{base}/api/presentations/slide",
                        json={"deck_id": deck_id, "index": i})
            if sr.status_code != 200:
                R.add(f"slide {i} generated", False, f"{sr.status_code}: {sr.text[:160]}")
                break
            slides.append(sr.json())
        else:
            R.add(f"all {len(slides)} slides generated",
                  len(slides) == len(titles),
                  f"in {time.time() - t0:.1f}s")
        outline = {"title": plan_data.get("title") or topic,
                   "subtitle": plan_data.get("subtitle") or "",
                   "slides": slides}
        # resume check: regenerating one slide must work standalone
        if deck_id and slides:
            again = c.post(f"{base}/api/presentations/slide",
                           json={"deck_id": deck_id, "index": 0})
            R.add("single-slide resume/retry", again.status_code == 200,
                  f"{again.status_code}")

    if outline is None:
        gen = c.post(f"{base}/api/presentations/generate",
                     json={"topic": topic, "num_slides": 6})
        if gen.status_code != 200:
            R.add("presentation generate (legacy)", False,
                  f"{gen.status_code}: {gen.text[:200]}")
            return
        outline = gen.json()
        R.add("presentation generate (legacy)", True,
              f"{len(outline.get('slides', []))} slides")

    pptx = c.post(f"{base}/api/presentations/pptx",
                  json={"outline": outline, "theme": "ocean"})
    if pptx.status_code != 200:
        R.add("pptx export", False, f"{pptx.status_code}: {pptx.text[:200]}")
    else:
        blob = pptx.content
        ok_zip = blob[:2] == b"PK"
        slide_count = -1
        try:
            with zipfile.ZipFile(io.BytesIO(blob)) as zf:
                slide_count = len([n for n in zf.namelist()
                                   if re.match(r"ppt/slides/slide\d+\.xml$", n)])
        except Exception as exc:  # noqa: BLE001
            R.add("pptx is a valid container", False, str(exc))
        R.add("pptx export", ok_zip and slide_count >= len(outline["slides"]),
              f"{len(blob)} bytes, {slide_count} slides")
        try:
            from pptx import Presentation

            prs = Presentation(io.BytesIO(blob))
            texts = "\n".join(
                sh.text_frame.text for s in prs.slides for sh in s.shapes
                if getattr(sh, "has_text_frame", False))
            R.add("pptx opens in python-pptx", True, f"{len(prs.slides)} slides")
            R.add("pptx contains slide text",
                  outline["slides"][0]["title"][:12].lower() in texts.lower(),
                  outline["slides"][0]["title"][:40])
        except ImportError:
            R.skip("pptx opens in python-pptx", "python-pptx not installed here")

    html = c.post(f"{base}/api/presentations/html",
                  json={"outline": outline, "theme": "ocean"})
    if html.status_code == 200:
        page = html.text
        R.add("HTML deck export",
              "<html" in page.lower() and outline["slides"][0]["title"][:10] in page,
              f"{len(page)} bytes")
    else:
        R.add("HTML deck export", False, f"{html.status_code}: {html.text[:200]}")


def test_images(c: httpx.Client, base: str) -> None:
    t0 = time.time()
    r = c.post(f"{base}/api/images/generate",
               json={"prompt": "a minimalist geometric logo on dark charcoal",
                     "width": 1024, "height": 1024})
    if r.status_code != 200:
        R.add("image request", False, f"{r.status_code}: {r.text[:200]}")
        return
    url = r.json().get("url")
    R.add("image request", bool(url), f"{time.time() - t0:.1f}s")
    if not url:
        return
    try:
        img = c.get(url)
    except Exception as exc:  # noqa: BLE001 — a network hiccup must not abort the run
        R.add("image bytes returned", False, f"{type(exc).__name__}: {exc}"[:160])
        return
    ctype = img.headers.get("content-type", "")
    R.add("image bytes returned",
          img.status_code == 200 and len(img.content) > 4000 and ctype.startswith("image/"),
          f"{img.status_code} {ctype} {len(img.content)} bytes in {time.time() - t0:.1f}s")


def test_voice(c: httpx.Client, base: str) -> None:
    voices = c.get(f"{base}/api/voice/voices")
    if voices.status_code == 200:
        data = voices.json()
        names = [v.get("id") for v in (data.get("voices") or [])]
        R.add("voice list", len(names) >= 3, ", ".join(names[:6]))
    else:
        R.add("voice list", False, f"{voices.status_code}")

    for phrase, label in (("Hi there, this is a quick voice test.", "short"),
                          ("Hello. I can help with homework, ideas, and "
                           "quick questions. What would you like to talk about?", "two-sentence")):
        t0 = time.time()
        r = c.post(f"{base}/api/voice/tts",
                   json={"text": phrase, "voice": "nova", "speed": 1.0})
        ttfb = time.time() - t0
        if r.status_code == 503:
            # Documented free-tier fallback: the browser takes over with a
            # neural system voice. That is a pass, but worth calling out.
            detail = r.json().get("detail", {}) if r.headers.get("content-type", "").startswith("application/json") else {}
            R.add(f"TTS ({label}) browser fallback offered",
                  bool(detail.get("browser_tts")),
                  f"503, retry_after={detail.get('retry_after')}")
            continue
        if r.status_code != 200:
            R.add(f"TTS ({label})", False, f"{r.status_code}: {r.text[:160]}")
            continue
        blob = r.content
        ctype = r.headers.get("content-type", "")
        looks_audio = blob[:3] == b"ID3" or blob[:2] == b"\xff\xfb" or blob[:4] == b"RIFF" or (
            ctype.startswith("audio/"))
        R.add(f"TTS ({label})", looks_audio and len(blob) > 2000,
              f"{len(blob)} bytes, {ctype}, ttfb {ttfb:.1f}s")
        if label == "short":
            R.add("TTS first sound fast", ttfb < 12, f"{ttfb:.1f}s")


def test_hd_voice(c: httpx.Client, base: str) -> None:
    """SOFT probe: the HD Edge voices are a free, unofficial service.

    A busy or blocked HD endpoint is *not* a failure — the call chain is
    HD → backup voice → browser voices. Everything here is a warning, and the
    'HD TTS' line always lands in the sticky report.
    """
    catalogue = c.get(f"{base}/api/voice/voices")
    if catalogue.status_code != 200:
        R.warn("HD TTS", f"voices list returned {catalogue.status_code}")
        return
    hd = (catalogue.json() or {}).get("hd") or {}
    voices = hd.get("voices") or []
    if not hd:
        R.warn("HD TTS", "no 'hd' block in /api/voice/voices")
        return
    uz = [v for v in voices if str(v.get("short", "")).startswith("uz-")]
    R.warn("HD TTS catalogue",
           f"enabled={hd.get('enabled')} voices={len(voices)} "
           f"languages={','.join(l.get('code', '') for l in (hd.get('languages') or [])[:4])} "
           f"uzbek={'yes' if uz else 'no'}")
    R.warn("HD TTS default voice", str(hd.get("default_voice") or "?"))

    t0 = time.time()
    try:
        r = c.post(f"{base}/api/voice/tts",
                   json={"text": "Salom, bu qisqa sinov.", "voice": "uz-UZ-SardorNeural",
                         "speed": 1.0, "provider": "auto"})
    except Exception as exc:  # noqa: BLE001 — never fail the run on a soft probe
        R.warn("HD TTS", f"request failed: {type(exc).__name__}")
        return
    took = time.time() - t0
    engine = r.headers.get("x-aether-tts", "")
    if r.status_code == 200:
        size = len(r.content)
        if engine == "hd":
            R.warn("HD TTS",
                   f"uz-UZ-SardorNeural answered with HD audio: {size} bytes in {took:.1f}s")
            R.warn("HD TTS size check",
                   f"{'PASS' if size > 5120 else 'small'} — {size} bytes "
                   f"({'above' if size > 5120 else 'below'} the 5 KB expectation)")
        else:
            R.warn("HD TTS",
                   f"HD busy — backup voice answered ({engine or 'legacy'}), "
                   f"{size} bytes in {took:.1f}s; the fallback chain held")
        return
    detail = {}
    if r.headers.get("content-type", "").startswith("application/json"):
        detail = (r.json() or {}).get("detail") or {}
    R.warn("HD TTS",
           f"unavailable right now ({r.status_code}, browser_tts="
           f"{bool(detail.get('browser_tts'))}) — the chain fell back cleanly")


def test_status(c: httpx.Client, base: str) -> None:
    r = c.get(f"{base}/api/ai/status")
    if r.status_code != 200:
        R.add("ai status", False, f"{r.status_code}")
        return
    data = r.json()
    R.add("ai status", bool(data.get("providers")), json.dumps(data)[:120])
    keyless = [p for p in data.get("providers", []) if p.get("kind") == "pollinations"]
    R.add("keyless provider present", bool(keyless),
          json.dumps(keyless)[:160] if keyless else "missing")
    if keyless and "queue" in keyless[0]:
        R.note(f"keyless queue: {json.dumps(keyless[0]['queue'])}")
    hd = data.get("hd_tts")
    if isinstance(hd, dict):
        R.warn("HD TTS status", json.dumps(hd))
    else:
        R.warn("HD TTS status", "not reported by /api/ai/status")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--base-url", required=True)
    ap.add_argument("--report", default="",
                    help="write the markdown report to this file")
    args = ap.parse_args()
    base = args.base_url.rstrip("/")

    print(f"== Aether keyless E2E ==\n{base}\n")
    with client() as c:
        test_static(c, base)
        if not test_auth(c, base):
            print("\nCannot continue without an account.", file=sys.stderr)
            return 2
        conv = c.post(f"{base}/api/conversations", json={"title": "E2E chat"})
        if conv.status_code not in (200, 201):
            R.add("create conversation", False, f"{conv.status_code} {conv.text[:160]}")
            conversation = None
        else:
            conversation = conv.json()["id"]
            R.add("create conversation", True, conversation)
        if conversation:
            test_chat(c, base, conversation)
            test_chat_controls(c, base, conversation)
        test_teach_and_memory(c, base)
        test_usage(c, base)
        test_presentations(c, base)
        test_images(c, base)
        test_voice(c, base)
        test_hd_voice(c, base)
        test_status(c, base)

    markdown = R.summary_markdown()
    print("\n" + markdown)
    if args.report:
        verdict = ("❌ " + ", ".join(n for n, _, _ in R.failed)) if R.failed else "✅ all green"
        with open(args.report, "w", encoding="utf-8") as fh:
            fh.write(f"Target: `{base}` — {verdict}\n\n{markdown}\n")
    summary_path = os.getenv("GITHUB_STEP_SUMMARY")
    if summary_path:
        with open(summary_path, "a", encoding="utf-8") as fh:
            fh.write("## Aether keyless E2E (preview)\n\n")
            fh.write(f"Target: `{base}`\n\n{R.summary_markdown()}\n")
    return 1 if R.failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
